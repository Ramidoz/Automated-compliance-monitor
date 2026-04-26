from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_CHARS = 200_000


class UnsupportedFileType(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    elif suffix in {".txt", ".md", ".html"}:
        text = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileType(f"Unsupported file type: {suffix}")
    return _normalize(text)[:MAX_CHARS]


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text)
    return "\n".join(parts)


def _normalize(text: str) -> str:
    lines = [line.rstrip() for line in text.splitlines()]
    cleaned: list[str] = []
    blank = 0
    for line in lines:
        if line.strip():
            cleaned.append(line)
            blank = 0
        else:
            blank += 1
            if blank <= 1:
                cleaned.append("")
    return "\n".join(cleaned).strip()
